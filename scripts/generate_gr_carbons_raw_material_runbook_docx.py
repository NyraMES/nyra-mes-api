"""Generate GR Carbons raw material setup + inward runbook as Word (.docx)."""

from __future__ import annotations

import sys
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Pt

OUT = ROOT / "docs" / "GR_Carbons_Raw_Material_Setup_and_Inward_Steps.docx"


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_para(doc: Document, text: str, bold: bool = False) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(11)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Number")


def main() -> int:
    doc = Document()
    title = doc.add_heading("GR Carbons — Raw Material Setup & Inward Steps", 0)
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    sub = doc.add_paragraph()
    sub.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    r = sub.add_run(f"Nyra MES Activated Carbon · Tenant GR_CARBONS · {date.today().isoformat()}")
    r.italic = True
    r.font.size = Pt(10)

    doc.add_paragraph()

    add_heading(doc, "1. Environment", 1)
    add_para(
        doc,
        "All data is stored in PostgreSQL (database nyra_mes_core). "
        "Local stack uses tenant GR_CARBONS and workspace GR_CARBONS_ACTIVATED_CARBON.",
    )
    add_heading(doc, "Start services (local)", 2)
    add_bullets(
        doc,
        [
            "API: nyra-mes-api\\scripts\\run_sam_carbons_api_8003.ps1 → http://127.0.0.1:8003",
            "UI: nyra-mes-ui → npm run dev:ac → http://localhost:5173",
            "Plant setup: http://localhost:5173/activated-carbon/masters",
            "Raw Material: http://localhost:5173/activated-carbon/raw-material",
            "API docs (workflow steps): http://127.0.0.1:8003/docs",
        ],
    )

    add_heading(doc, "2. Clean reset (empty plant)", 1)
    add_para(doc, "Run when you need to wipe transactional data and start again:")
    add_para(doc, "PowerShell (venv active, MES_CORE_DB_PASSWORD in .env):", bold=True)
    add_para(doc, "  cd nyra-mes-api")
    add_para(doc, "  .\\scripts\\reset_gr_carbons_client.ps1")
    add_bullets(
        doc,
        [
            "Removes all inward orders, lots, production orders, and KG triples for GR_CARBONS.",
            "Reinstalls template masters (raw/finished SKUs, kiln/resources, demo customer).",
            "Suppliers count goes to 0 — add suppliers again in Masters.",
            "Restart API and hard-refresh the browser after reset.",
        ],
    )

    add_heading(doc, "3. Super Admin — Plant setup (Masters)", 1)
    add_para(doc, "Do once after each reset, in this order.")

    add_heading(doc, "3.1 Suppliers", 2)
    add_numbered(
        doc,
        [
            "Open Masters → Suppliers tab.",
            "Add at least one supplier (dummy names are fine for trial).",
            "Save each record; note SUP-AC-* codes if shown.",
        ],
    )

    add_heading(doc, "3.2 Materials & steps", 2)
    add_numbered(
        doc,
        [
            "Open Masters → Materials & steps (Super Administrator).",
            "Choose existing raw material (e.g. RAW-COCO-SHELL) OR + Add new raw material…",
            "Finished product: top bar on the right — “Choose raw material → then becomes → Finished product”. "
            "Editable only when adding a new raw material (output product name). "
            "Existing selection shows finished SKU and routing read-only (e.g. AC-GRAN-8x30 · RT-AC-8x30-V1).",
            "Flow diagram — fixed steps (always present): Required Raw Material Inward → "
            "optional inbound prep steps → Required Raw Material QC.",
            "Inbound prep: for GR notebook (lab, bags, crusher, RM outward), leave empty on first setup "
            "unless you want extra Crushing/Screening steps with triplet parameters.",
            "Shop floor (after release): apply standard template or add Kiln Feed → Activation → "
            "Screening → QC; open each step to confirm parameters and mandatory flags.",
            "Save — material type must be ACTIVE.",
        ],
    )

    add_heading(doc, "3.3 Advanced settings (Ontology & UI)", 2)
    add_numbered(
        doc,
        [
            "Open Masters → Advanced settings.",
            "Under Raw material QC parameters, enable fields GR uses: moisture, foreign matter, "
            "stones, unburnt, -20 mesh fines, operator remarks; disable unused fields.",
            "Keep ontology predicates from the dropdown (triplet URIs are not changed in code).",
            "Save configuration.",
        ],
    )

    add_heading(doc, "3.4 Machines & kilns (optional check)", 2)
    add_bullets(
        doc,
        [
            "Confirm RES-KILN-01, screen, lab resources from template SQL.",
            "Add crusher resource if you will reference it on crusher performance.",
        ],
    )

    add_heading(doc, "4. GR process map (your written flow)", 1)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Notebook step"
    hdr[1].text = "Nyra concept"
    hdr[2].text = "Where to run today"
    rows = [
        (
            "Received at gate (supplier, DC, weight, value)",
            "Raw Material Inward + receipt triples on ORD-AC-*",
            "UI: Raw Material → Inward tab",
        ),
        (
            "Lab (moisture, dust, AD, approve/deny, remarks)",
            "Initial QC (inbound workflow)",
            "API: POST …/raw-material-initial-qc (Swagger on port 8003)",
        ),
        (
            "RMS inward — select inward, bag + weight",
            "Inward bags",
            "API: inward bag endpoints; partial UI: Bag traceability menu",
        ),
        (
            "Crusher performance (rotary, meshes, moisture, dust)",
            "Crusher output",
            "API: weight-to-crusher, crusher-output",
        ),
        (
            "RM outward / grades (-20 Rot A/B, etc.)",
            "Grade separation",
            "API: raw-material-grade-separation",
        ),
        (
            "Final inspect before production",
            "Raw Material QC",
            "UI: Raw Material → QC tab (fields from Advanced settings)",
        ),
        (
            "Release to activation path",
            "Release to production",
            "UI: Raw Material → Release tab → LOT-AC-* / PO-AC-*",
        ),
    ]
    for a, b, c in rows:
        row = table.add_row().cells
        row[0].text = a
        row[1].text = b
        row[2].text = c

    doc.add_paragraph()

    add_heading(doc, "5. Operator steps — Raw material inward (first receipt)", 1)

    add_heading(doc, "5.1 Gate receipt", 2)
    add_numbered(
        doc,
        [
            "Go to Raw Material → Inward tab.",
            "Select supplier, raw material, quantity (kg).",
            "Save — first order after reset is typically ORD-AC-001.",
        ],
    )

    add_heading(doc, "5.2 Lab initial QC", 2)
    add_numbered(
        doc,
        [
            "Open http://127.0.0.1:8003/docs.",
            "Find POST /api/activated-carbon/raw-material-initial-qc.",
            "Use customer_order_id from the inward order, moisture_pct, dust_pct, ad_value, qc_status PASS or FAIL.",
        ],
    )

    add_heading(doc, "5.3 Bags, crusher, grade separation", 2)
    add_para(
        doc,
        "Complete in order using the related activated-carbon endpoints in Swagger "
        "(bags → weight to crusher → crusher output → grade separation). "
        "Triplet subjects remain urn:nyra:ac:inward:{order_code} and related bag/crusher URIs.",
    )

    add_heading(doc, "5.4 Raw Material QC", 2)
    add_numbered(
        doc,
        [
            "Raw Material → QC tab.",
            "Select ORD-AC-001.",
            "Enter enabled QC fields; set PASS and remarks; save.",
        ],
    )

    add_heading(doc, "5.5 Release to production", 2)
    add_numbered(
        doc,
        [
            "Raw Material → Release tab.",
            "Select QC-passed inward; release to production.",
            "Note PO-AC-* and LOT-AC-* for feasibility / kiln / execution (next documents).",
        ],
    )

    add_heading(doc, "6. Verification", 1)
    add_bullets(
        doc,
        [
            "KG Triples page: confirm staged triples for the inward order after saves.",
            "Health check: http://127.0.0.1:8003/api/health should show status ok.",
        ],
    )

    add_heading(doc, "7. Notes", 1)
    add_bullets(
        doc,
        [
            "Triplet ontology (predicate URIs and subject patterns) is fixed in application code; "
            "configuration changes which QC fields appear and which process steps are saved on the raw material type.",
            "UI tabs for lab / bags / crusher / RM outward can be added so operators need not use Swagger.",
            "Cognito / testbed: set custom:tenant_code to GR_CARBONS when deploying.",
        ],
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"Wrote {OUT}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
