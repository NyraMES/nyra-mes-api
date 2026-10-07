"""GR Carbons process-flow presentation (editable PowerPoint).

Usage (repo root, venv):
    python scripts/generate_gr_carbons_process_flow_presentation.py
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "GR_Carbons_NyraMES_Process_Flow_Presentation.pptx"

NAVY = RGBColor(0x0D, 0x3A, 0x6E)
NAVY_DARK = RGBColor(0x0A, 0x2C, 0x54)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x0F, 0x17, 0x2A)
MUTED = RGBColor(0x47, 0x55, 0x69)
GOLD = RGBColor(0xC9, 0xA2, 0x27)
CARD = RGBColor(0xF1, 0xF5, 0xF9)
ACCENT = RGBColor(0x1D, 0x4E, 0x89)

W = Inches(13.333)
H = Inches(7.5)


def _set_run(run, text, size=18, bold=False, color=INK, font="Calibri"):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def _fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def _rect(slide, l, t, w, h, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    _fill(shape, color)
    shape.adjustments[0] = 0.08
    return shape


def _box_text(shape, lines, size=14, bold=False, color=WHITE, align=PP_ALIGN.CENTER):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.clear()
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        _set_run(run, line, size=size, bold=bold if i == 0 else False, color=color)


def _textbox(slide, l, t, w, h, text, size=16, bold=False, color=INK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    _set_run(run, text, size=size, bold=bold, color=color)
    return box


def _bullets(slide, l, t, w, h, items, size=18, color=INK):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        p.space_after = Pt(8)
        run = p.add_run()
        _set_run(run, "•  " + item, size=size, color=color)
    return box


def _header(slide, title, subtitle=None):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, Inches(1.15))
    _fill(bar, NAVY)
    gold = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(1.15), W, Inches(0.06))
    _fill(gold, GOLD)
    _textbox(slide, Inches(0.45), Inches(0.22), Inches(12.4), Inches(0.5), title, 26, True, WHITE)
    if subtitle:
        _textbox(slide, Inches(0.45), Inches(0.68), Inches(12.4), Inches(0.35), subtitle, 14, False, RGBColor(0xD6, 0xE4, 0xF5))
    foot = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(7.22), W, Inches(0.28))
    _fill(foot, NAVY_DARK)
    _textbox(
        slide,
        Inches(0.45),
        Inches(7.24),
        Inches(12.4),
        Inches(0.24),
        "Nyra MES  ·  GR Carbons Activated Carbon  ·  Process flow  ·  Editable working copy",
        11,
        False,
        RGBColor(0xCB, 0xD5, 0xE1),
    )


def _blank(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    _fill(bg, WHITE)
    return slide


def _flow_row(slide, top, items, box_w=1.7, gap=0.12):
    n = len(items)
    total = n * box_w + (n - 1) * gap
    left = (13.333 - total) / 2
    shapes = []
    for i, (label, color) in enumerate(items):
        x = Inches(left + i * (box_w + gap))
        shape = _rect(slide, x, Inches(top), Inches(box_w), Inches(0.85), color)
        lines = label.split("\n")
        _box_text(shape, lines, size=12 if len(lines) > 1 else 13, bold=True, color=WHITE)
        shapes.append(shape)
        if i < n - 1:
            arr = slide.shapes.add_shape(
                MSO_SHAPE.RIGHT_ARROW,
                Inches(left + i * (box_w + gap) + box_w - 0.02),
                Inches(top + 0.32),
                Inches(0.18),
                Inches(0.22),
            )
            _fill(arr, GOLD)
    return shapes


def main() -> int:
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H

    # 1 Title
    s = _blank(prs)
    hero = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    _fill(hero, NAVY)
    gold = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(5.85), W, Inches(0.08))
    _fill(gold, GOLD)
    _textbox(s, Inches(0.7), Inches(1.8), Inches(12), Inches(0.4), "GR CARBONS  ·  NYRA MES", 16, True, GOLD)
    _textbox(s, Inches(0.7), Inches(2.25), Inches(12), Inches(1.4), "Activated Carbon\nProcess Flow", 44, True, WHITE)
    _textbox(
        s,
        Inches(0.7),
        Inches(4.7),
        Inches(12),
        Inches(0.8),
        "From supplier inward to kiln, post-activation, customer PO and dispatch.\nWorking copy — edit freely before the meeting.",
        18,
        False,
        RGBColor(0xE2, 0xE8, 0xF0),
    )

    # 2 Agenda
    s = _blank(prs)
    _header(s, "Agenda", "What we will walk through")
    _bullets(
        s,
        Inches(0.7),
        Inches(1.55),
        Inches(12),
        Inches(5.3),
        [
            "How the plant is modelled in Nyra MES",
            "End-to-end material flow (one picture)",
            "Raw Material — gate, lab, RMS, crusher, release",
            "Execution — Activation (kiln) then Post-activation",
            "Customer PO, In Stock, and Dispatch",
            "Who uses which screens (Admin vs shop floor)",
        ],
        22,
    )

    # 3 How the app is organised
    s = _blank(prs)
    _header(s, "How the application is organised", "Left menu in Nyra MES — Activated Carbon")
    modules = [
        ("Overview", "Dashboard\nKiln telemetry", NAVY),
        ("Production", "Raw Material\nExecution\nCustomer PO\nDispatch\nReports", ACCENT),
        ("Admin", "Users\nCustomer\nSupplier", RGBColor(0x0F, 0x76, 0x6E)),
        ("Masters", "Materials & steps\nMachines & kilns\nRouting\nInbound settings", RGBColor(0x7C, 0x3A, 0xED)),
    ]
    for i, (title, body, color) in enumerate(modules):
        x = Inches(0.55 + i * 3.15)
        card = _rect(s, x, Inches(1.7), Inches(2.95), Inches(4.7), CARD)
        head = _rect(s, x, Inches(1.7), Inches(2.95), Inches(0.7), color)
        _box_text(head, [title], 18, True, WHITE)
        tf = s.shapes.add_textbox(x + Inches(0.18), Inches(2.55), Inches(2.6), Inches(3.5)).text_frame
        tf.word_wrap = True
        for j, line in enumerate(body.split("\n")):
            p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            p.space_after = Pt(10)
            run = p.add_run()
            _set_run(run, "•  " + line, 16, False, INK)

    # 4 End to end
    s = _blank(prs)
    _header(s, "End-to-end plant flow", "One chain — from the gate to the customer")
    _flow_row(
        s,
        1.7,
        [
            ("Gate\nreceipt", NAVY),
            ("Inbound\nprep", ACCENT),
            ("Release to\nproduction", RGBColor(0x0F, 0x76, 0x6E)),
            ("Kiln\nactivation", RGBColor(0xB4, 0x53, 0x09)),
            ("Post-\nactivation", RGBColor(0x7C, 0x3A, 0xED)),
            ("In Stock &\nDispatch", RGBColor(0xBE, 0x12, 0x3C)),
        ],
        box_w=1.95,
        gap=0.18,
    )
    notes = [
        ("Shop identity", "Plant users work with the inward order ORD-AC-* until release. They think in bags (RI / RO / KA / DS), not internal lot numbers."),
        ("After release", "Execution starts at kiln feed. Load on Kiln A or Kiln B. Remaining activation and post-activation steps are tabs on Execution."),
        ("Stock to ship", "Bags marked In Stock after destoning or post-activation become available on Customer PO and Dispatch."),
    ]
    for i, (h, body) in enumerate(notes):
        y = Inches(2.85 + i * 1.35)
        card = _rect(s, Inches(0.55), y, Inches(12.2), Inches(1.2), CARD)
        _textbox(s, Inches(0.8), y + Inches(0.12), Inches(11.7), Inches(0.35), h, 16, True, NAVY)
        _textbox(s, Inches(0.8), y + Inches(0.48), Inches(11.7), Inches(0.6), body, 15, False, INK)

    # 5 Roles
    s = _blank(prs)
    _header(s, "Who uses which screens", "Plant roles — keep this slide if you want to assign logins in the meeting")
    rows = [
        ("Admin", "Users, Customer, Supplier. Dashboard only besides Admin."),
        ("RMS", "Raw Material — gate, lab, bags, crusher, outward, release."),
        ("Activation", "Execution — kiln feed, kiln output, destoning, kiln parameters."),
        ("Post Activation", "Execution — screening through finished goods."),
        ("Quality", "Quality tabs on kiln feed, kiln output, destoning, post-activation."),
        ("Manager", "Plant screens plus Customer PO, Dispatch, Reports."),
        ("nyraservices", "Everything, including Masters and Routing Builder."),
    ]
    for i, (role, desc) in enumerate(rows):
        y = Inches(1.5 + i * 0.75)
        chip = _rect(s, Inches(0.55), y, Inches(2.6), Inches(0.62), NAVY if i != 6 else GOLD)
        _box_text(chip, [role], 14, True, WHITE if i != 6 else NAVY)
        body = _rect(s, Inches(3.3), y, Inches(9.45), Inches(0.62), CARD)
        _textbox(s, Inches(3.5), y + Inches(0.14), Inches(9.1), Inches(0.4), desc, 16, False, INK)

    # 6 RM overview
    s = _blank(prs)
    _header(s, "Raw Material — four tabs", "Screen: Raw Material")
    tabs = [
        ("1 · Gate receipt", "Supplier inward at the dock.\nDC number, supplier kg, our kg."),
        ("2 · Inbound prep", "Lab → RMS inward → Crusher → RMS outward."),
        ("3 · Release", "Release the inward order to production.\nThen go to Execution."),
        ("4 · Reports", "Read-only. Process dropdown,\ncharts, Excel. No capture."),
    ]
    for i, (title, body) in enumerate(tabs):
        x = Inches(0.45 + i * 3.2)
        card = _rect(s, x, Inches(1.7), Inches(3.05), Inches(4.7), CARD)
        head = _rect(s, x, Inches(1.7), Inches(3.05), Inches(1.0), NAVY)
        _box_text(head, [title], 16, True, WHITE)
        _textbox(s, x + Inches(0.18), Inches(2.95), Inches(2.7), Inches(3.1), body, 16, False, INK)

    # 7 Gate
    s = _blank(prs)
    _header(s, "Gate receipt", "First capture — identity of the inward is ORD-AC-*")
    _bullets(
        s,
        Inches(0.7),
        Inches(1.55),
        Inches(7.2),
        Inches(5.2),
        [
            "Supplier (from Admin → Supplier)",
            "Raw material type (Coconut Shell)",
            "DC number — required",
            "Supplier weight (KG) — required",
            "Our weight (KG) — this is received quantity",
            "System creates the inward order ORD-AC-*",
            "A production order and lot are created in the background for trace — not shown to the operator",
        ],
        20,
    )
    card = _rect(s, Inches(8.2), Inches(1.7), Inches(4.5), Inches(4.6), NAVY)
    _box_text(
        card,
        [
            "Operator sees",
            "",
            "ORD-AC-1",
            "",
            "Not PO-AC-*",
            "Not LOT-AC-*",
            "",
            "“Lot” on the floor =",
            "RMS outward bags",
        ],
        18,
        True,
        WHITE,
    )

    # 8 Inbound prep
    s = _blank(prs)
    _header(s, "Inbound prep — four steps after the gate", "Must be completed in this order on the same inward")
    steps = [
        ("Lab", "IN-LAB", "Moisture %, dust %, AD.\nApprove or Deny.\nDeny needs remarks and stops this inward only."),
        ("RMS inward", "IN-RMS-BAG", "Bags RI_BAG_001…\nTotal bag KG must equal our weight from the gate."),
        ("Crusher", "IN-CRUSHER", "Feed RMS bags one by one.\nGrade, dust %, moisture %, rotary."),
        ("RMS outward", "IN-RM-OUT", "Output bags RO_BAG_001…\nRotary A or B.\nGrade + Unburnt + Stones."),
    ]
    for i, (name, code, body) in enumerate(steps):
        x = Inches(0.4 + i * 3.22)
        card = _rect(s, x, Inches(1.65), Inches(3.08), Inches(5.15), CARD)
        head = _rect(s, x, Inches(1.65), Inches(3.08), Inches(1.25), NAVY if i < 3 else RGBColor(0x0F, 0x76, 0x6E))
        _box_text(head, [f"{i + 1}. {name}", code], 16, True, WHITE)
        _textbox(s, x + Inches(0.16), Inches(3.1), Inches(2.76), Inches(3.4), body, 15, False, INK)

    # 9 Rotary / grades
    s = _blank(prs)
    _header(s, "RMS outward — rotary grades", "Configured in Masters → Inbound settings (nyraservices)")
    _bullets(
        s,
        Inches(0.7),
        Inches(1.55),
        Inches(12),
        Inches(5.3),
        [
            "Choose Rotary A or Rotary B for crusher output bags",
            "That rotary needs a grade value, plus Unburnt and Stones",
            "Rotary A typically: -20 1st Stage, Grade 1st stage",
            "Rotary B typically: -20 2nd Stage, Grade 2nd stage",
            "Kiln feed loads only rotary grades — Unburnt and Stones stay off the kiln",
            "Release shows the RO_BAG_* details from this step",
        ],
        20,
    )

    # 10 Release
    s = _blank(prs)
    _header(s, "Release to production", "Raw Material tab 3 — then Execution")
    _bullets(
        s,
        Inches(0.7),
        Inches(1.55),
        Inches(12),
        Inches(5.3),
        [
            "Completing RMS outward auto-marks the inward so release can proceed (no separate RM QC screen)",
            "Release allocates charcoal to production",
            "Operator then opens Execution — kiln feed is first",
            "Feasibility gate is not used on the floor (it redirects to Execution)",
            "Lot routing is not a separate screen — Load on the kiln creates the work tabs",
        ],
        20,
    )

    # 11 Activation
    s = _blank(prs)
    _header(s, "Execution — Activation", "Sidebar: Execution → Activation")
    acts = [
        ("Kiln feed", "Load RO_BAG_* onto Kiln A or Kiln B.\nLive load graph.\nQuality on feed bags."),
        ("Kiln output", "Output bags KA_DDMMYY_##### or KB_DDMMYY_#####.\nOutput quality."),
        ("Destoning", "Creates DS_DDMMYY_#####.\nDestination can be In Stock."),
        ("Kiln parameters", "Hourly readings with timestamp.\nNot lot-specific — a plant reading."),
        ("Reports", "Read-only activation report.\nProcess flow pane is hidden."),
    ]
    for i, (title, body) in enumerate(acts):
        y = Inches(1.5 + i * 1.05)
        chip = _rect(s, Inches(0.55), y, Inches(3.1), Inches(0.9), NAVY)
        _box_text(chip, [title], 16, True, WHITE)
        body_s = _rect(s, Inches(3.85), y, Inches(8.9), Inches(0.9), CARD)
        _textbox(s, Inches(4.05), y + Inches(0.15), Inches(8.5), Inches(0.65), body, 16, False, INK)

    # 12 Post activation
    s = _blank(prs)
    _header(s, "Execution — Post-activation", "Sidebar: Execution → Post activation")
    _flow_row(
        s,
        1.7,
        [
            ("Screening\nSC_", ACCENT),
            ("De-dusting\nDD_", RGBColor(0x0F, 0x76, 0x6E)),
            ("Crushing\nCR_", NAVY),
            ("De-mag\nDM_", RGBColor(0x7C, 0x3A, 0xED)),
            ("Blending\nBL_", RGBColor(0xB4, 0x53, 0x09)),
            ("Finished\ngoods FG_", RGBColor(0xBE, 0x12, 0x3C)),
        ],
        box_w=1.95,
        gap=0.18,
    )
    _bullets(
        s,
        Inches(0.7),
        Inches(2.85),
        Inches(12),
        Inches(4.0),
        [
            "Each station has live load and output bags dated DDMMYY",
            "De-dusting and de-magnetisation output grade is Exklin only",
            "Destination In Stock makes those bags available for Customer PO",
            "Quality tab on post-activation; Reports tab is read-only",
            "Process flow stays on the right pane — it is removed when Reports is open",
        ],
        20,
    )

    # 13 IDs
    s = _blank(prs)
    _header(s, "Numbers the plant will see", "Each key has its own PostgreSQL sequence")
    ids = [
        ("ORD-AC-*", "Inward / shop identity"),
        ("RI_BAG_*", "RMS inward bags"),
        ("RO_BAG_*", "RMS outward bags"),
        ("KA_ / KB_", "Kiln A / Kiln B output"),
        ("DS_*", "Destoning bags"),
        ("SC_ DD_ CR_\nDM_ BL_ FG_", "Post-activation bags"),
        ("CUST_PO_*", "Customer sales PO"),
        ("DISP-AC-*", "Dispatch document"),
    ]
    for i, (code, meaning) in enumerate(ids):
        col = i % 4
        row = i // 4
        x = Inches(0.5 + col * 3.2)
        y = Inches(1.65 + row * 2.5)
        card = _rect(s, x, y, Inches(3.0), Inches(2.25), CARD)
        head = _rect(s, x, y, Inches(3.0), Inches(1.05), NAVY)
        _box_text(head, code.split("\n"), 16, True, WHITE)
        _textbox(s, x + Inches(0.15), y + Inches(1.2), Inches(2.7), Inches(0.85), meaning, 15, False, INK)

    # 14 Customer PO / Dispatch
    s = _blank(prs)
    _header(s, "Customer PO and Dispatch", "Finished carbon leaving the plant")
    left = _rect(s, Inches(0.5), Inches(1.6), Inches(6.0), Inches(5.2), CARD)
    right = _rect(s, Inches(6.8), Inches(1.6), Inches(6.0), Inches(5.2), CARD)
    lh = _rect(s, Inches(0.5), Inches(1.6), Inches(6.0), Inches(0.7), NAVY)
    rh = _rect(s, Inches(6.8), Inches(1.6), Inches(6.0), Inches(0.7), RGBColor(0xBE, 0x12, 0x3C))
    _box_text(lh, ["Customer PO"], 18, True, WHITE)
    _box_text(rh, ["Dispatch"], 18, True, WHITE)
    _bullets(
        s,
        Inches(0.75),
        Inches(2.5),
        Inches(5.5),
        Inches(4.0),
        [
            "Create CUST_PO_* against a customer",
            "Available stock = destoning / post-activation bags with destination In Stock",
            "Not leftover FG metadata",
        ],
        17,
    )
    _bullets(
        s,
        Inches(7.05),
        Inches(2.5),
        Inches(5.5),
        Inches(4.0),
        [
            "Ship against a customer PO — shown at the top",
            "Pick In Stock bags (or oldest first)",
            "Dispatch identity DISP-AC-*",
        ],
        17,
    )

    # 15 Dashboard
    s = _blank(prs)
    _header(s, "Dashboard and Reports", "Management view vs shop-floor capture")
    _bullets(
        s,
        Inches(0.7),
        Inches(1.55),
        Inches(12),
        Inches(5.3),
        [
            "Dashboard tabs: Overview, Raw Material, Activation, Post-activation, Machine utilization",
            "High-level KPIs: charcoal at dock, GC, plant yield, kiln yield, Exklin with / without stone",
            "Machine counts come from Masters → Machines & kilns",
            "Section Reports (RM / Activation / Post-activation) are read-only — charts and Excel",
            "Reports do not write shop-floor data and do not change the process flow",
        ],
        20,
    )

    # 16 Admin vs Masters
    s = _blank(prs)
    _header(s, "Admin vs Masters", "Who is allowed to change plant setup")
    left = _rect(s, Inches(0.5), Inches(1.6), Inches(6.0), Inches(5.2), CARD)
    right = _rect(s, Inches(6.8), Inches(1.6), Inches(6.0), Inches(5.2), CARD)
    lh = _rect(s, Inches(0.5), Inches(1.6), Inches(6.0), Inches(0.7), RGBColor(0x0F, 0x76, 0x6E))
    rh = _rect(s, Inches(6.8), Inches(1.6), Inches(6.0), Inches(0.7), GOLD)
    _box_text(lh, ["Admin  —  plant administrator"], 16, True, WHITE)
    _box_text(rh, ["Masters  —  nyraservices only"], 16, True, NAVY)
    _bullets(
        s,
        Inches(0.75),
        Inches(2.5),
        Inches(5.5),
        Inches(4.0),
        [
            "Users (create plant logins)",
            "Customer (who you sell to)",
            "Supplier (who delivers charcoal)",
        ],
        18,
    )
    _bullets(
        s,
        Inches(7.05),
        Inches(2.5),
        Inches(5.5),
        Inches(4.0),
        [
            "Materials & process steps",
            "Machines, kilns, routing",
            "Kiln reading fields",
            "Inbound rotary settings",
        ],
        18,
    )

    # 17 Units / rules
    s = _blank(prs)
    _header(s, "Shop-floor rules worth stating in the meeting", "These are live in the application today")
    _bullets(
        s,
        Inches(0.7),
        Inches(1.55),
        Inches(12),
        Inches(5.3),
        [
            "Weights are KG. Moisture and dust are %. AD is a lab reading, not a weight",
            "RMS inward bags have no grade; outward bags carry the rotary split as grade",
            "Lab deny stops that inward only — other inwards continue",
            "Do not click capture buttons twice; the screen stays locked until the save finishes",
            "Process flow is on the right of Execution; it is hidden on Reports",
        ],
        20,
    )

    # 18 Suggested walkthrough
    s = _blank(prs)
    _header(s, "Suggested live walkthrough", "One receipt through to a dispatch")
    _bullets(
        s,
        Inches(0.7),
        Inches(1.55),
        Inches(12),
        Inches(5.3),
        [
            "Admin: confirm supplier and customer exist",
            "Raw Material: gate receipt → lab approve → RMS bags → crusher → RMS outward",
            "Release to production",
            "Execution Activation: load kiln → output bags → destoning (optional In Stock)",
            "Execution Post-activation: screening → … → finished goods / In Stock",
            "Customer PO against In Stock → Dispatch against that PO",
            "Dashboard / Reports to show the same kg in management view",
        ],
        20,
    )

    # 19 Close
    s = _blank(prs)
    hero = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    _fill(hero, NAVY)
    gold = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(5.85), W, Inches(0.08))
    _fill(gold, GOLD)
    _textbox(s, Inches(0.7), Inches(2.2), Inches(12), Inches(1.2), "One chain. One plant.", 40, True, WHITE)
    _textbox(
        s,
        Inches(0.7),
        Inches(3.6),
        Inches(12),
        Inches(1.4),
        "Supplier inward  →  kiln  →  post-activation  →  In Stock  →  customer.\nNyra MES keeps the shop-floor capture and the trace together.",
        20,
        False,
        RGBColor(0xE2, 0xE8, 0xF0),
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"Wrote {OUT}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
