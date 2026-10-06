"""Generate GR Carbons trial process-flow deck (python-pptx)."""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "GR_Carbons_NyraMES_Activated_Carbon_Process_Flow_Triplets.pptx"


def _add_title_slide(prs: Presentation, title: str, subtitle: str) -> None:
    layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = title
    slide.placeholders[1].text = subtitle


def _add_section(prs: Presentation, title: str, bullets: list[str]) -> None:
    layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = title
    tf = slide.placeholders[1].text_frame
    tf.clear()
    for i, line in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.level = 0
        p.font.size = Pt(20)


def _add_two_column(prs: Presentation, title: str, left_title: str, left: list[str], right_title: str, right: list[str]) -> None:
    layout = prs.slide_layouts[5]  # Title Only — we'll use text boxes
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = title
    box_l = slide.shapes.add_textbox(Inches(0.5), Inches(1.4), Inches(4.5), Inches(5.5))
    box_r = slide.shapes.add_textbox(Inches(5.2), Inches(1.4), Inches(4.5), Inches(5.5))
    for box, head, items in ((box_l, left_title, left), (box_r, right_title, right)):
        tf = box.text_frame
        tf.text = head
        tf.paragraphs[0].font.bold = True
        tf.paragraphs[0].font.size = Pt(18)
        for item in items:
            p = tf.add_paragraph()
            p.text = item
            p.level = 0
            p.font.size = Pt(16)


def main() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    _add_title_slide(
        prs,
        "Nyra MES — Activated Carbon Process Flow",
        "1-month trial for GR Carbons\nSupplier inward → production lot → shop floor → finished goods → dispatch",
    )

    _add_section(
        prs,
        "Scope of the trial",
        [
            "Raw material: inward receipt, QC, release to production (PO + LOT)",
            "Production gate: feasibility, kiln feed, lot routing",
            "Shop-floor execution with step-wise data capture",
            "Finished-goods in stock and customer dispatch",
            "Dashboard, traceability, kiln parameters & telemetry (as enabled)",
        ],
    )

    _add_section(
        prs,
        "End-to-end flow (bird's eye view)",
        [
            "SUPPLIER INWARD → RAW MATERIAL QC → RELEASE TO PRODUCTION",
            "↓",
            "FEASIBILITY GATE → KILN FEED → LOT ROUTING → EXECUTION",
            "↓",
            "FINISHED GOODS IN STOCK → DISPATCH",
            "",
            "Inward order: ORD-AC-*  ·  Production lot: LOT-AC-*  ·  Work orders from routing",
            "After release, operators and planners work on LOT codes, not inward orders.",
        ],
    )

    _add_section(
        prs,
        "Phase 1 — Raw material (3 screens)",
        [
            "1. Inward — supplier, raw material type, quantity received (KG)",
            "2. Raw material QC — lab/yard tests; Pass / Fail / Hold",
            "3. Release to production — allocate KG; creates production order + LOT-AC-*",
            "Release may create one lot or multiple lots (one per bag grade when grade splits exist).",
            "All weights in KG unless noted as % on QC forms.",
        ],
    )

    _add_two_column(
        prs,
        "Data captured — raw material",
        "Inward",
        [
            "Supplier",
            "Raw material type",
            "Quantity received (KG)",
            "Optional / extensible: DC number, supplier weight (KG), value (INR)",
        ],
        "Raw material QC (configurable)",
        [
            "Moisture, ash, foreign matter (%)",
            "Stones, unburnt, -20 mesh fines (KG)",
            "QC status & operator remarks",
            "Stored on inward order; feasibility checks QC passed",
        ],
    )

    _add_section(
        prs,
        "Optional inbound depth (platform capability)",
        [
            "Same inward order can also store: initial QC (moisture, dust, AD)",
            "Bag split & reconcile → weight to crusher",
            "Crusher: rotary A/B; mesh scrap +2 … -14 (KG); moisture & dust (%)",
            "Six bag-grade splits → mass balance & multi-lot release",
            "Primary trial UI: Inward + QC + Release — extended steps aligned with GR SOP in workshop",
        ],
    )

    _add_section(
        prs,
        "Inbound — triplet ontology (order-scoped)",
        [
            "Subject: urn:nyra:ac:inward:{ORD-AC-code} (before release)",
            "Inward receipt — supplies / captures: DC, supplier weight, received KG, value INR",
            "Initial QC — testedBy → initialQc; moisture, dust, AD parameters",
            "Bags — containedIn; crusher & grade separation — captures mesh scrap & grade KG",
            "Raw Material QC tab — testedBy → rawMaterialQc; stones / unburnt / -20 KG predicates",
        ],
    )

    _add_section(
        prs,
        "Phase 2 — Production gate",
        [
            "Feasibility Gate — select LOT-AC-*; checks material (QC), kiln availability, resources",
            "Kiln Feed — charge activation kiln; lot IN_PROGRESS",
            "Lot Routing — generate work orders from raw-material routing template (Masters)",
            "Routing defines pyrolysis, activation, screening, grade separation, FG QC, etc.",
        ],
    )

    _add_section(
        prs,
        "Triplet ontology — what it is",
        [
            "Semantic trace layer: kg.triple_store_staging (PostgreSQL) + MES operational tables",
            "Each save can stage subject · predicate · object (+ literal parameters)",
            "Inward steps use subject: urn:nyra:ac:inward:{ORD-AC-*}",
            "After release, production & execution use: urn:nyra:ac:lot:{LOT-AC-*}",
            "Same data operators enter in UI — triples enable traceability, analytics, Ask Nyra",
        ],
    )

    _add_two_column(
        prs,
        "Phase 3 — Shop-floor execution (operational data)",
        "Every routing step (Execution screen)",
        [
            "Start · Complete · Hold on each work order",
            "Actual output (KG) · Scrap (KG) · Downtime (min)",
            "Operator / labor pool · Machine when required",
            "Routing parameters (temps, steam, time, …) on complete",
            "Stored: exec.operation_execution + parameter readings",
        ],
        "Dedicated step panels (Execution)",
        [
            "SCREENING — mesh; good / overs / fines (KG)",
            "GRADE_SEPARATION — grade code; separated KG; yield",
            "QC — iodine, ash, moisture, bulk density; disposition",
            "Release to in-stock moves lot to FG stock when pipeline complete",
        ],
    )

    _add_section(
        prs,
        "Execution — triplet ontology (lot-scoped)",
        [
            "Subject for all rows below: urn:nyra:ac:lot:{LOT-AC-code}",
            "",
            "SCREENING (on save in Execution)",
            "  • screenedAt → urn:nyra:ac:screening:{execution_code}",
            "  • captures: SCREEN_MESH_SIZE, GOOD_KG, OVERS_KG, FINES_KG",
            "",
            "GRADE SEPARATION (on save in Execution)",
            "  • captures → urn:nyra:ac:gradeSeparation:{execution_code}",
            "  • captures: GRADE_CODE, SEPARATED_KG, GRADE_YIELD_PCT, OVERS_GRADE_KG, FINES_GRADE_KG",
        ],
    )

    _add_section(
        prs,
        "Execution — triplet ontology (continued)",
        [
            "QUALITY TESTING (on save — Release to in-stock)",
            "  • testedBy → urn:nyra:ac:qualityCheck:{QC-AC-code} (ontology)",
            "  • Lab values in qms.quality_check metadata (iodine, ash, moisture, bulk density)",
            "  • hasIodineValue staged on lot when iodine entered",
            "",
            "DISPATCH (when shipping)",
            "  • dispatchedTo — lot linked to customer shipment",
            "",
            "Pyrolysis / activation / kiln hourly params: operational executions + lot metadata",
            "  (extend triple parameters in trial where GR needs explicit KG predicates)",
        ],
    )

    _add_section(
        prs,
        "Pre-execution triples (same lot subject)",
        [
            "Feasibility Gate — feasibilityChecked → urn:nyra:ac:feasibility:{lot}",
            "Kiln Feed — assignedTo → urn:nyra:ac:kiln:{kiln_resource_code}",
            "Lot Routing — generates → urn:nyra:ac:routing:{lot}",
            "Links planner actions to the same LOT-AC-* trace chain as Execution",
        ],
    )

    _add_section(
        prs,
        "Phase 4 — Stock & dispatch",
        [
            "Release to in-stock — when routing complete, screening recorded, QC passed",
            "Lot status & FG stock updated — visible on Dispatch",
            "Customer PO — order against finished product; stock check",
            "Dispatch — ship from in-stock by grade; traceability ORD → LOT → QC → dispatch",
        ],
    )

    _add_section(
        prs,
        "Supporting capabilities",
        [
            "Operational Intelligence Hub — KPIs, pipeline, kiln strip",
            "Kiln parameters — hourly / zone readings per lot",
            "Kiln telemetry — live Modbus path (testbed)",
            "Masters — suppliers, materials, kilns, routing template, QC/quality fields",
            "Reports, traceability, bag tracing (capabilities per tenant)",
        ],
    )

    _add_section(
        prs,
        "Configured for GR Carbons",
        [
            "Raw material types and fixed routing per feedstock",
            "QC & finished-goods quality fields (Masters → Ontology)",
            "Plant resources — kilns, labor pools, tools",
            "Grade names & mesh bands — workshop with GR during trial",
            "Trial validates your SOP against this model, not a generic demo only",
        ],
    )

    _add_section(
        prs,
        "Trial access",
        [
            "URL: https://nyraservices-testbed.net/mes/",
            "Module: Activated Carbon · secure login (GR tenant / workspace)",
            "Architecture: PostgreSQL nyra_mes_core, REST API, audit trail",
            "Suggested pilot: one inward → QC → release → one lot through feasibility, kiln, routing, execution, in-stock",
        ],
    )

    _add_section(
        prs,
        "Workshop — what else should we capture?",
        [
            "Mandatory inward fields (DC, weighbridge, commercial value)?",
            "Crusher / bag / mesh scrap on every receipt — UI priority?",
            "Actual routing steps vs plant (kiln → activation → screening → QC)?",
            "Mandatory lab tests on raw vs finished carbon?",
            "COA, certificates, photos? PLC / weighbridge / LIMS later?",
        ],
    )

    _add_section(
        prs,
        "Summary",
        [
            "One chain: Supplier → QC → Lot → Kiln → Execute → Stock → Dispatch",
            "Lot-centric after release (LOT-AC-*)",
            "Operational data + triplet ontology on key execution steps (screening, grades, QC)",
            "Configurable for GR Carbons grades, tests, and routing",
            "Next: joint testbed walkthrough + configuration workshop",
        ],
    )

    prs.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
